from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from accounts.catalog import ROLE_STAFF
from accounts.models import Property
from accounts.property_excel import PROPERTY_COLUMNS, LOCATION_COLUMNS, SHEET_LOCATIONS, SHEET_PROPERTIES, SHEET_REFERENCE
from accounts.services import create_user, initialize_access_control, set_individual_permissions
from cases.models import Case
from crm.models import Customer
from locations.models import Location, LocationType
from pipeline.models import PipelineStage
from quotes.models import Quote, QuoteStatus


class PropertyExcelTests(TestCase):
    def setUp(self):
        initialize_access_control()
        self.superuser = User.objects.create_superuser(
            username="root",
            email="root@example.com",
            password="super-secret-pass",
        )
        self.assigned_staff = create_user(
            username="staff.lee",
            password="staff-pass-123",
            first_name="Lee",
            last_name="Staff",
            role=ROLE_STAFF,
        )
        self.other_staff = create_user(
            username="staff.other",
            password="staff-pass-123",
            first_name="Omar",
            last_name="Other",
            role=ROLE_STAFF,
        )
        self.kl = Property.objects.create(
            name="Hotel KL",
            code="KLX",
            region="Kuala Lumpur",
            city="Kuala Lumpur",
            rate="150.00",
            country="Malaysia",
            currency="MYR",
        )
        self.penang = Property.objects.create(
            name="Hotel Penang",
            code="PENX",
            region="Penang",
            city="George Town",
            rate="120.00",
            country="Malaysia",
            currency="MYR",
        )
        self.kl.staff.add(self.assigned_staff)
        self.room = Location.objects.create(
            property=self.kl,
            name="Room 101",
            code="101",
            location_type=LocationType.ROOM,
            capacity=2,
            status="available",
            base_rate="150.00",
            currency="MYR",
        )
        Location.objects.create(
            property=self.penang,
            name="Villa 01",
            code="V01",
            location_type=LocationType.VILLA,
            capacity=4,
            status="available",
        )
        self.customer = Customer.objects.create(name="John Smith", email="john@example.com", property=self.kl)
        self.case = Case.objects.create(
            sequence=1,
            case_number="CASE-9001",
            customer=self.customer,
            property=self.kl,
            location=self.room,
            title="Weekend stay",
            stage=PipelineStage.WON,
        )
        self.quote = Quote.objects.create(
            sequence=1,
            quote_number="Q-9001",
            case=self.case,
            amount="300.00",
            currency="MYR",
            status=QuoteStatus.PENDING,
        )

    def _login(self, username, password):
        response = self.client.post(
            "/api/auth/login/",
            data={"username": username, "password": password},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        return response

    def _xlsx(self, content, name="upload.xlsx"):
        return SimpleUploadedFile(name, content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def _workbook(self, properties=None, locations=None, property_headers=None, location_headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_PROPERTIES
        sheet.append(property_headers or PROPERTY_COLUMNS)
        for row in properties or []:
            sheet.append(row)
        loc = workbook.create_sheet(SHEET_LOCATIONS)
        loc.append(location_headers or LOCATION_COLUMNS)
        for row in locations or []:
            loc.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_properties_export_is_xlsx_and_scoped(self):
        self._login("staff.lee", "staff-pass-123")
        response = self.client.get("/api/properties/export/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertTrue(response["Content-Disposition"].endswith('filename="properties_export.xlsx"'))
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn(SHEET_PROPERTIES, workbook.sheetnames)
        rows = list(workbook[SHEET_PROPERTIES].iter_rows(min_row=2, values_only=True))
        names = {row[1] for row in rows}
        self.assertEqual(names, {"Hotel KL"})
        self.assertEqual(rows[0][0], self.kl.pk)
        self.assertNotIn("password", [cell.value for cell in workbook[SHEET_PROPERTIES][1]])

    def test_locations_export_and_portfolio_sheets(self):
        self._login("root", "super-secret-pass")
        locations = self.client.get("/api/locations/export/")
        self.assertEqual(locations.status_code, 200)
        loc_book = load_workbook(BytesIO(locations.content))
        self.assertIn(SHEET_LOCATIONS, loc_book.sheetnames)
        loc_rows = list(loc_book[SHEET_LOCATIONS].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(loc_rows), 2)

        scoped = self.client.get("/api/locations/export/", {"property_id": self.kl.pk})
        scoped_book = load_workbook(BytesIO(scoped.content))
        self.assertEqual(len(list(scoped_book[SHEET_LOCATIONS].iter_rows(min_row=2, values_only=True))), 1)

        portfolio = self.client.get("/api/properties/portfolio-export/")
        self.assertEqual(portfolio.status_code, 200)
        book = load_workbook(BytesIO(portfolio.content))
        self.assertEqual(set(book.sheetnames), {SHEET_PROPERTIES, SHEET_LOCATIONS, SHEET_REFERENCE})
        self.assertEqual(len(list(book[SHEET_PROPERTIES].iter_rows(min_row=2, values_only=True))), 2)

    def test_assigned_user_cannot_export_foreign_locations(self):
        self._login("staff.lee", "staff-pass-123")
        response = self.client.get("/api/locations/export/", {"property_id": self.penang.pk})
        self.assertEqual(response.status_code, 403)

    def test_template_download(self):
        self._login("root", "super-secret-pass")
        response = self.client.get("/api/properties/import-template/")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertIn(SHEET_PROPERTIES, book.sheetnames)
        self.assertIn(SHEET_LOCATIONS, book.sheetnames)
        self.assertIn(SHEET_REFERENCE, book.sheetnames)

    def test_update_existing_property_and_location_preserves_ids_and_relations(self):
        self._login("root", "super-secret-pass")
        content = self._workbook(
            properties=[
                [
                    self.kl.pk,
                    "Hotel KL City",
                    "KLX",
                    "hotel",
                    "Malaysia",
                    "Kuala Lumpur",
                    "Kuala Lumpur",
                    "Jalan Sultan Ismail",
                    "",
                    "50250",
                    "+603-1111-1111",
                    "kl@demo.ibilik.local",
                    "",
                    "Updated hotel",
                    "MYR",
                    175.5,
                    "active",
                    "15:00",
                    "12:00",
                    "Asia/Kuala_Lumpur",
                    "",
                    "",
                ]
            ],
            locations=[
                [
                    self.room.pk,
                    self.kl.pk,
                    "Hotel KL City",
                    "KLX",
                    "Deluxe Room 101",
                    "101",
                    "room",
                    3,
                    2,
                    1,
                    "1",
                    "queen",
                    1,
                    "Updated room",
                    180,
                    "MYR",
                    "available",
                    "Yes",
                ]
            ],
        )
        preview = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(content), "mode": "portfolio"})
        self.assertEqual(preview.status_code, 200, preview.content)
        self.assertTrue(preview.json()["valid"])
        self.assertEqual(preview.json()["summary"]["properties_update"], 1)
        self.assertEqual(preview.json()["summary"]["locations_update"], 1)

        committed = self.client.post("/api/properties/import/", data={"file": self._xlsx(content), "mode": "portfolio"})
        self.assertEqual(committed.status_code, 200, committed.content)
        self.kl.refresh_from_db()
        self.room.refresh_from_db()
        self.case.refresh_from_db()
        self.quote.refresh_from_db()
        self.assertEqual(self.kl.name, "Hotel KL City")
        self.assertEqual(self.kl.pk,  self.kl.id)
        self.assertEqual(str(self.kl.rate), "175.50")
        self.assertEqual(self.room.name, "Deluxe Room 101")
        self.assertEqual(self.room.capacity, 3)
        self.assertEqual(self.case.property_id, self.kl.pk)
        self.assertEqual(self.case.location_id, self.room.pk)
        self.assertEqual(self.quote.case_id, self.case.pk)

    def test_create_property_and_location_from_portfolio(self):
        self._login("root", "super-secret-pass")
        content = self._workbook(
            properties=[
                [
                    "",
                    "Hotel JB",
                    "JHBX",
                    "business_hotel",
                    "Malaysia",
                    "Johor",
                    "Johor Bahru",
                    "Johor Bahru, Johor",
                    "",
                    "80000",
                    "",
                    "jb@demo.ibilik.local",
                    "",
                    "New hotel",
                    "MYR",
                    160,
                    "active",
                    "14:00",
                    "12:00",
                    "Asia/Kuala_Lumpur",
                    "",
                    "",
                ]
            ],
            locations=[
                [
                    "",
                    "",
                    "Hotel JB",
                    "JHBX",
                    "Room 201",
                    "201",
                    "suite",
                    2,
                    2,
                    0,
                    "2",
                    "king",
                    1,
                    "Suite",
                    220,
                    "MYR",
                    "available",
                    "Yes",
                ]
            ],
        )
        response = self.client.post("/api/properties/import/", data={"file": self._xlsx(content), "mode": "portfolio"})
        self.assertEqual(response.status_code, 200, response.content)
        created = Property.objects.get(code="JHBX")
        location = Location.objects.get(property=created, code="201")
        self.assertEqual(created.city, "Johor Bahru")
        self.assertEqual(location.location_type, "suite")
        self.assertNotEqual(created.pk, self.kl.pk)

    def test_invalid_file_type_and_missing_sheet(self):
        self._login("root", "super-secret-pass")
        bad_type = self.client.post(
            "/api/properties/import/preview/",
            data={"file": SimpleUploadedFile("hotels.xls", b"\xd0\xcf\x11\xe0", content_type="application/vnd.ms-excel"), "mode": "properties"},
        )
        self.assertEqual(bad_type.status_code, 400)
        self.assertEqual(bad_type.json()["code"], "invalid_file_type")

        workbook = Workbook()
        workbook.active.title = "Wrong"
        buffer = BytesIO()
        workbook.save(buffer)
        missing = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(buffer.getvalue()), "mode": "properties"})
        self.assertEqual(missing.status_code, 400)
        self.assertEqual(missing.json()["code"], "missing_sheet")

    def test_validation_errors_do_not_commit(self):
        self._login("root", "super-secret-pass")
        original_name = self.kl.name
        content = self._workbook(
            properties=[
                [
                    self.kl.pk,
                    "Hacked Name",
                    "KLX",
                    "spaceship",
                    "Malaysia",
                    "Kuala Lumpur",
                    "Kuala Lumpur",
                    "",
                    "",
                    "",
                    "",
                    "not-an-email",
                    "",
                    "",
                    "!!!",
                    "abc",
                    "closed",
                    "soon",
                    "later",
                    "Not/AZone",
                    "",
                    "",
                ]
            ]
        )
        preview = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(content), "mode": "properties"})
        self.assertEqual(preview.status_code, 200)
        self.assertFalse(preview.json()["valid"])
        self.assertGreater(preview.json()["summary"]["properties_errors"], 0)

        committed = self.client.post("/api/properties/import/", data={"file": self._xlsx(content), "mode": "properties"})
        self.assertEqual(committed.status_code, 422)
        self.kl.refresh_from_db()
        self.assertEqual(self.kl.name, original_name)

    def test_duplicate_ids_and_location_codes_and_capacity(self):
        self._login("root", "super-secret-pass")
        content = self._workbook(
            properties=[
                [self.kl.pk, "Hotel KL", "KLX", "hotel", "Malaysia", "Kuala Lumpur", "Kuala Lumpur", "", "", "", "", "", "", "", "MYR", 150, "active", "15:00", "12:00", "Asia/Kuala_Lumpur", "", ""],
                [self.kl.pk, "Hotel KL Copy", "KL2", "hotel", "Malaysia", "Kuala Lumpur", "Kuala Lumpur", "", "", "", "", "", "", "", "MYR", 150, "active", "15:00", "12:00", "Asia/Kuala_Lumpur", "", ""],
            ],
            locations=[
                ["", self.kl.pk, "Hotel KL", "KLX", "Room A", "101", "room", 0, "", "", "", "", "", "", "", "MYR", "available", "Yes"],
                ["", self.kl.pk, "Hotel KL", "KLX", "Room B", "101", "room", 2, "", "", "", "", "", "", "", "MYR", "available", "Yes"],
            ],
        )
        preview = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(content), "mode": "portfolio"})
        self.assertEqual(preview.status_code, 200)
        messages = " ".join(item["message"] for item in preview.json()["issues"])
        self.assertIn("Duplicate Property ID", messages)
        self.assertIn("Duplicate Location Code", messages)
        self.assertIn("Capacity", messages)

    def test_unauthorized_property_import_is_forbidden_and_atomic(self):
        self._login("staff.lee", "staff-pass-123")
        content = self._workbook(
            properties=[
                [self.penang.pk, "Hacked Penang", "PENX", "hotel", "Malaysia", "Penang", "George Town", "", "", "", "", "", "", "", "MYR", 99, "active", "15:00", "12:00", "Asia/Kuala_Lumpur", "", ""]
            ]
        )
        preview = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(content), "mode": "properties"})
        self.assertEqual(preview.status_code, 200)
        self.assertFalse(preview.json()["valid"])
        self.assertTrue(any("access" in item["message"].lower() for item in preview.json()["issues"]))
        self.penang.refresh_from_db()
        self.assertEqual(self.penang.name, "Hotel Penang")

        viewer = create_user(username="viewer.only", password="staff-pass-123", role=ROLE_STAFF)
        set_individual_permissions(viewer, ["properties.view_property"])
        viewer.groups.clear()
        self.kl.staff.add(viewer)
        self.client.logout()
        self._login("viewer.only", "staff-pass-123")
        blocked = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(content), "mode": "properties"})
        self.assertEqual(blocked.status_code, 403)

    def test_missing_required_columns(self):
        self._login("root", "super-secret-pass")
        content = self._workbook(properties=[["Hotel"]], property_headers=["Wrong"])
        preview = self.client.post("/api/properties/import/preview/", data={"file": self._xlsx(content), "mode": "properties"})
        self.assertEqual(preview.status_code, 400)
        self.assertEqual(preview.json()["code"], "missing_columns")

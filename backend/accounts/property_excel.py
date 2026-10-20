from __future__ import annotations

from datetime import datetime, time
from decimal import Decimal
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from accounts.authorization import can_access_property, user_has_permission
from accounts.models import Property, PropertyStatus, PropertyType
from accounts.property_services import create_property, properties_for_user, update_property
from locations.models import BedType, Location, LocationStatus, LocationType
from locations.services import create_location, update_location

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

SHEET_PROPERTIES = "Properties"
SHEET_LOCATIONS = "Locations"
SHEET_REFERENCE = "Reference Data"

PROPERTY_COLUMNS = [
    "Property ID",
    "Property Name",
    "Property Code",
    "Property Type",
    "Country",
    "Region",
    "City",
    "Address",
    "Address Line 2",
    "Postal Code",
    "Phone",
    "Email",
    "Website",
    "Description",
    "Currency",
    "Base Rate",
    "Status",
    "Check-in Time",
    "Check-out Time",
    "Timezone",
    "Staff Count",
    "Location Count",
]

LOCATION_COLUMNS = [
    "Location ID",
    "Property ID",
    "Property Name",
    "Property Code",
    "Location Name",
    "Location Code",
    "Location Type",
    "Capacity",
    "Max Adults",
    "Max Children",
    "Floor",
    "Bed Type",
    "Bed Count",
    "Description",
    "Base Rate",
    "Currency",
    "Status",
    "Is Bookable",
]

PROPERTY_EXPORT_ONLY = {"Staff Count", "Location Count"}
LOCATION_EXPORT_ONLY = {"Property Name"}

HEADER_FILL = PatternFill("solid", fgColor="111827")
HEADER_FONT = Font(bold=True, color="FFFFFF")


class ExcelImportError(ValueError):
    def __init__(self, message: str, *, code: str = "invalid_excel"):
        super().__init__(message)
        self.code = code


def excel_http_response(content: bytes, filename: str):
    from django.http import HttpResponse

    response = HttpResponse(content, content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _style_sheet(worksheet, widths: list[int]) -> None:
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for cell in worksheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _write_headers(worksheet, headers: list[str]) -> None:
    worksheet.append(headers)


def _as_time_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value)


def _as_number(value):
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    return value


def build_reference_rows() -> list[list[str]]:
    rows = [["Category", "Value", "Label"]]
    for item in PropertyType:
        rows.append(["Property Type", item.value, item.label])
    for item in PropertyStatus:
        rows.append(["Property Status", item.value, item.label])
    for item in LocationType:
        rows.append(["Location Type", item.value, item.label])
    for item in LocationStatus:
        rows.append(["Location Status", item.value, item.label])
    for item in BedType:
        rows.append(["Bed Type", item.value, item.label])
    rows.append(["Currency", "MYR", "Malaysian Ringgit"])
    return rows


def property_export_row(property_obj: Property) -> list:
    staff_count = getattr(property_obj, "staff_count", None)
    if staff_count is None:
        staff_count = property_obj.staff.count()
    location_count = getattr(property_obj, "location_count", None)
    if location_count is None:
        location_count = property_obj.locations.count()
    return [
        property_obj.pk,
        property_obj.name,
        property_obj.code,
        property_obj.property_type,
        property_obj.country,
        property_obj.region,
        property_obj.city,
        property_obj.address_line1,
        property_obj.address_line2,
        property_obj.postal_code,
        property_obj.phone,
        property_obj.email,
        property_obj.website,
        property_obj.description,
        property_obj.currency,
        _as_number(property_obj.rate),
        property_obj.status,
        _as_time_text(property_obj.check_in_time),
        _as_time_text(property_obj.check_out_time),
        property_obj.timezone,
        staff_count,
        location_count,
    ]


def location_export_row(location: Location) -> list:
    return [
        location.pk,
        location.property_id,
        location.property.name,
        location.property.code,
        location.name,
        location.code,
        location.location_type,
        location.capacity,
        location.max_adults,
        location.max_children,
        location.floor,
        location.bed_type,
        location.bed_count,
        location.description,
        _as_number(location.base_rate),
        location.currency or location.property.currency,
        location.status,
        "Yes" if location.is_bookable else "No",
    ]


def _apply_rate_formats(worksheet, header_map: dict[str, int]) -> None:
    rate_index = header_map.get("Base Rate")
    if not rate_index:
        return
    for row in worksheet.iter_rows(min_row=2, min_col=rate_index, max_col=rate_index):
        for cell in row:
            if isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = "0.00"


def build_workbook(*, properties=None, locations=None, include_reference: bool = False, example: bool = False) -> bytes:
    workbook = Workbook()
    properties_sheet = workbook.active
    properties_sheet.title = SHEET_PROPERTIES
    _write_headers(properties_sheet, PROPERTY_COLUMNS)
    for property_obj in properties or []:
        properties_sheet.append(property_export_row(property_obj))
    if example and not properties:
        properties_sheet.append(
            [
                "",
                "iBilik Example Hotel",
                "EX",
                "hotel",
                "Malaysia",
                "Kuala Lumpur",
                "Kuala Lumpur",
                "Jalan Sultan Ismail, Bukit Bintang",
                "",
                "50250",
                "+603-0000-0000",
                "reservations@demo.ibilik.local",
                "https://demo.ibilik.local",
                "Fictional demo hotel for import.",
                "MYR",
                200.00,
                "active",
                "15:00",
                "12:00",
                "Asia/Kuala_Lumpur",
                "",
                "",
            ]
        )
    _style_sheet(properties_sheet, [14, 36, 14, 18, 14, 18, 18, 36, 22, 14, 18, 32, 28, 28, 12, 12, 12, 14, 14, 22, 12, 14])
    _apply_rate_formats(properties_sheet, {name: index for index, name in enumerate(PROPERTY_COLUMNS, start=1)})

    locations_sheet = workbook.create_sheet(SHEET_LOCATIONS)
    _write_headers(locations_sheet, LOCATION_COLUMNS)
    for location in locations or []:
        locations_sheet.append(location_export_row(location))
    if example and not locations:
        locations_sheet.append(
            [
                "",
                "",
                "iBilik Example Hotel",
                "EX",
                "Deluxe Room 101",
                "101",
                "room",
                2,
                2,
                0,
                "1",
                "queen",
                1,
                "City-view room",
                200.00,
                "MYR",
                "available",
                "Yes",
            ]
        )
    _style_sheet(locations_sheet, [12, 12, 28, 14, 28, 14, 16, 12, 12, 12, 10, 12, 12, 28, 12, 12, 14, 14])
    _apply_rate_formats(locations_sheet, {name: index for index, name in enumerate(LOCATION_COLUMNS, start=1)})

    if include_reference:
        reference = workbook.create_sheet(SHEET_REFERENCE)
        for row in build_reference_rows():
            reference.append(row)
        _style_sheet(reference, [20, 22, 28])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_properties_workbook(user) -> bytes:
    return build_workbook(properties=list(properties_for_user(user)), include_reference=False)


def export_locations_workbook(user, *, property_id: int | None = None) -> bytes:
    queryset = Location.objects.select_related("property").order_by("property__name", "name", "id")
    if property_id is not None:
        property_obj = Property.objects.filter(pk=property_id).first()
        if property_obj is None:
            raise ExcelImportError("Property not found.", code="not_found")
        if not can_access_property(user, property_obj):
            raise PermissionError("You do not have access to this property.")
        queryset = queryset.filter(property=property_obj)
    else:
        accessible_ids = properties_for_user(user).values_list("id", flat=True)
        queryset = queryset.filter(property_id__in=accessible_ids)
    return build_workbook(locations=list(queryset), include_reference=False)


def export_portfolio_workbook(user, *, property_id: int | None = None) -> bytes:
    properties = properties_for_user(user)
    if property_id is not None:
        properties = properties.filter(pk=property_id)
        if not properties.exists():
            property_obj = Property.objects.filter(pk=property_id).first()
            if property_obj is None:
                raise ExcelImportError("Property not found.", code="not_found")
            raise PermissionError("You do not have access to this property.")
    property_list = list(properties)
    locations = list(
        Location.objects.filter(property__in=property_list).select_related("property").order_by("property__name", "name", "id")
    )
    return build_workbook(properties=property_list, locations=locations, include_reference=True)


def export_template_workbook() -> bytes:
    return build_workbook(include_reference=True, example=True)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M") if value.hour or value.minute else value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _read_headers(worksheet) -> dict[str, int]:
    headers = {}
    for index, cell in enumerate(worksheet[1], start=1):
        name = _cell_text(cell.value)
        if name:
            headers[name] = index
    return headers


def _row_values(worksheet, row_index: int, headers: dict[str, int]) -> dict[str, object]:
    row = [cell.value for cell in worksheet[row_index]]
    values = {}
    for name, index in headers.items():
        values[name] = row[index - 1] if index - 1 < len(row) else None
    return values


def _normalize_header(name: str) -> str:
    return " ".join(name.strip().lower().replace("_", " ").split())


def _header_aliases() -> dict[str, str]:
    aliases = {}
    for name in PROPERTY_COLUMNS + LOCATION_COLUMNS:
        aliases[_normalize_header(name)] = name
    aliases["rate"] = "Base Rate"
    aliases["base rate"] = "Base Rate"
    aliases["type"] = "Property Type"
    aliases["address line 1"] = "Address"
    aliases["address line1"] = "Address"
    return aliases


def _canonicalize_headers(headers: dict[str, int]) -> tuple[dict[str, int], list[str]]:
    aliases = _header_aliases()
    canonical = {}
    unknown = []
    for name, index in headers.items():
        mapped = aliases.get(_normalize_header(name))
        if mapped:
            canonical[mapped] = index
        else:
            unknown.append(name)
    return canonical, unknown


def load_upload(uploaded) -> bytes:
    if uploaded is None:
        raise ExcelImportError("An Excel file is required.", code="missing_file")
    name = getattr(uploaded, "name", "") or ""
    if name.lower().endswith(".xls") and not name.lower().endswith(".xlsx"):
        raise ExcelImportError("Excel 97-2003 .xls files are not supported. Upload an .xlsx file.", code="invalid_file_type")
    if name and not name.lower().endswith(".xlsx"):
        raise ExcelImportError("Only .xlsx files are supported.", code="invalid_file_type")
    content = uploaded.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise ExcelImportError("The Excel file is too large.", code="file_too_large")
    if content[:2] != b"PK":
        raise ExcelImportError("Only .xlsx files are supported.", code="invalid_file_type")
    return content


def _open_workbook(content: bytes):
    try:
        return load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise ExcelImportError("The Excel file could not be read.", code="invalid_workbook") from exc


def _parse_int(value, *, field: str, required: bool = False, minimum: int | None = None, allow_zero: bool = False):
    if _is_blank(value):
        if required:
            raise ValueError(f"{field} is required.")
        return None
    if isinstance(value, bool):
        raise ValueError(f"{field} must be a valid integer.")
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(f"{field} must be a valid integer.")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field} must be a valid integer.") from exc
    if minimum is not None and number < minimum:
        raise ValueError(f"{field} must be at least {minimum}.")
    if not allow_zero and number < 1 and field.endswith("ID"):
        raise ValueError(f"{field} must be a positive integer.")
    return number


def _parse_choice(value, choices, *, field: str, required: bool = False, default=None):
    if _is_blank(value):
        if required:
            raise ValueError(f"{field} is required.")
        return default
    text = _cell_text(value)
    by_value = {item.value.lower(): item.value for item in choices}
    by_label = {item.label.lower(): item.value for item in choices}
    key = text.lower()
    if key in by_value:
        return by_value[key]
    if key in by_label:
        return by_label[key]
    raise ValueError(f"{field} is not supported.")


def _parse_bool(value, *, field: str, default=None):
    if _is_blank(value):
        return default
    if isinstance(value, bool):
        return value
    text = _cell_text(value).lower()
    if text in {"1", "true", "yes", "y", "on"}:
        return True
    if text in {"0", "false", "no", "n", "off"}:
        return False
    raise ValueError(f"{field} must be Yes or No.")


def _parse_time(value, *, field: str):
    if _is_blank(value):
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value)
            if isinstance(parsed, datetime):
                return parsed.strftime("%H:%M")
            if isinstance(parsed, time):
                return parsed.strftime("%H:%M")
        except Exception as exc:
            raise ValueError(f"{field} must be a valid time.") from exc
    text = _cell_text(value)
    try:
        time.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"{field} must be a valid time.") from exc
    return text


def _issue(sheet: str, row: int, column: str, message: str, level: str = "error") -> dict:
    return {"sheet": sheet, "row": row, "column": column, "level": level, "message": message}


def _sheet_rows(workbook, sheet_name: str, expected_columns: list[str], required_columns: list[str]):
    if sheet_name not in workbook.sheetnames:
        raise ExcelImportError(f'The "{sheet_name}" sheet is missing.', code="missing_sheet")
    worksheet = workbook[sheet_name]
    headers, unknown = _canonicalize_headers(_read_headers(worksheet))
    missing = [name for name in required_columns if name not in headers]
    if missing:
        raise ExcelImportError(
            f'The "{sheet_name}" sheet is missing required columns: {", ".join(missing)}.',
            code="missing_columns",
        )
    rows = []
    for row_index in range(2, worksheet.max_row + 1):
        values = _row_values(worksheet, row_index, headers)
        if all(_is_blank(value) for value in values.values()):
            continue
        rows.append((row_index, values))
    return rows, unknown


def preview_import(user, uploaded, *, mode: str = "portfolio", locked_property_id: int | None = None) -> dict:
    content = load_upload(uploaded)
    workbook = _open_workbook(content)
    if mode not in {"properties", "locations", "portfolio"}:
        raise ExcelImportError("Import mode is not supported.", code="invalid_mode")

    issues: list[dict] = []
    property_rows = []
    location_rows = []
    unknown_columns = []

    if mode in {"properties", "portfolio"}:
        required = ["Property Name", "Region", "City", "Base Rate"]
        rows, unknown = _sheet_rows(workbook, SHEET_PROPERTIES, PROPERTY_COLUMNS, required)
        unknown_columns.extend((SHEET_PROPERTIES, name) for name in unknown)
        property_rows = _validate_property_rows(user, rows, issues, locked_property_id=locked_property_id)
    elif SHEET_PROPERTIES in workbook.sheetnames:
        issues.append(_issue(SHEET_PROPERTIES, 1, "", "Properties sheet was ignored because this import is locations-only.", "warning"))

    pending_codes = {
        row["values"].get("code", "").upper()
        for row in property_rows
        if row["action"] == "create" and row["values"].get("code")
    }

    if mode in {"locations", "portfolio"}:
        if SHEET_LOCATIONS not in workbook.sheetnames:
            if mode == "portfolio":
                raise ExcelImportError(f'The "{SHEET_LOCATIONS}" sheet is missing.', code="missing_sheet")
        else:
            required = ["Location Name", "Location Code", "Location Type", "Capacity", "Status"]
            rows, unknown = _sheet_rows(workbook, SHEET_LOCATIONS, LOCATION_COLUMNS, required)
            unknown_columns.extend((SHEET_LOCATIONS, name) for name in unknown)
            location_rows = _validate_location_rows(
                user,
                rows,
                issues,
                property_rows=property_rows,
                pending_codes=pending_codes,
                locked_property_id=locked_property_id,
            )
    elif SHEET_LOCATIONS in workbook.sheetnames:
        issues.append(_issue(SHEET_LOCATIONS, 1, "", "Locations sheet was ignored because this import is properties-only.", "warning"))

    for sheet, column in unknown_columns:
        issues.append(_issue(sheet, 1, column, f'Unknown column "{column}" was ignored.', "warning"))

    summary = {
        "properties_create": sum(1 for row in property_rows if row["action"] == "create" and not row["errors"]),
        "properties_update": sum(1 for row in property_rows if row["action"] == "update" and not row["errors"]),
        "properties_errors": sum(1 for row in property_rows if row["errors"]),
        "locations_create": sum(1 for row in location_rows if row["action"] == "create" and not row["errors"]),
        "locations_update": sum(1 for row in location_rows if row["action"] == "update" and not row["errors"]),
        "locations_errors": sum(1 for row in location_rows if row["errors"]),
        "warnings": sum(1 for item in issues if item["level"] == "warning"),
        "errors": sum(1 for item in issues if item["level"] == "error"),
    }
    valid = summary["errors"] == 0 and summary["properties_errors"] == 0 and summary["locations_errors"] == 0
    if valid and not property_rows and not location_rows:
        valid = False
        issues.append(_issue(SHEET_PROPERTIES, 1, "", "The workbook does not contain any importable rows.", "error"))
        summary["errors"] += 1

    return {
        "resource": "property_import_preview",
        "valid": valid,
        "mode": mode,
        "summary": summary,
        "properties": property_rows,
        "locations": location_rows,
        "issues": issues,
    }


def _validate_property_rows(user, rows, issues, *, locked_property_id=None):
    results = []
    seen_ids = {}
    seen_codes = {}
    can_add = user_has_permission(user, "properties.add_property")
    can_change = user_has_permission(user, "properties.change_property")

    for row_index, raw in rows:
        errors = []
        warnings = []
        values = {}
        action = "create"
        property_obj = None

        try:
            property_id = _parse_int(raw.get("Property ID"), field="Property ID")
        except ValueError as exc:
            property_id = None
            errors.append(str(exc))
            issues.append(_issue(SHEET_PROPERTIES, row_index, "Property ID", str(exc)))

        if locked_property_id and property_id not in (None, locked_property_id):
            message = "This import can only update the selected property."
            errors.append(message)
            issues.append(_issue(SHEET_PROPERTIES, row_index, "Property ID", message))

        if property_id:
            action = "update"
            if property_id in seen_ids:
                message = "Duplicate Property ID in this workbook."
                errors.append(message)
                issues.append(_issue(SHEET_PROPERTIES, row_index, "Property ID", message))
            seen_ids[property_id] = row_index
            property_obj = Property.objects.filter(pk=property_id).first()
            if property_obj is None:
                message = "Property ID does not exist."
                errors.append(message)
                issues.append(_issue(SHEET_PROPERTIES, row_index, "Property ID", message))
            else:
                try:
                    if not can_access_property(user, property_obj):
                        raise PermissionError("You do not have access to this property.")
                except PermissionError as exc:
                    errors.append(str(exc))
                    issues.append(_issue(SHEET_PROPERTIES, row_index, "Property ID", str(exc)))
                if not can_change:
                    message = "You do not have permission to update properties."
                    errors.append(message)
                    issues.append(_issue(SHEET_PROPERTIES, row_index, "Property ID", message))

        code = _cell_text(raw.get("Property Code")).upper()
        if code:
            if code in seen_codes:
                message = "Duplicate Property Code in this workbook."
                errors.append(message)
                issues.append(_issue(SHEET_PROPERTIES, row_index, "Property Code", message))
            seen_codes[code] = row_index
            clash = Property.objects.filter(code=code)
            if property_obj:
                clash = clash.exclude(pk=property_obj.pk)
            if clash.exists():
                existing = clash.first()
                if property_id is None and can_access_property(user, existing):
                    action = "update"
                    property_obj = existing
                    property_id = existing.pk
                    warnings.append("Matched existing property by Property Code.")
                    issues.append(_issue(SHEET_PROPERTIES, row_index, "Property Code", warnings[-1], "warning"))
                    if not can_change:
                        message = "You do not have permission to update properties."
                        errors.append(message)
                        issues.append(_issue(SHEET_PROPERTIES, row_index, "Property Code", message))
                else:
                    message = "A property with this code already exists."
                    errors.append(message)
                    issues.append(_issue(SHEET_PROPERTIES, row_index, "Property Code", message))

        if action == "create" and not can_add:
            message = "You do not have permission to create properties."
            errors.append(message)
            issues.append(_issue(SHEET_PROPERTIES, row_index, "Property Name", message))

        def add_error(column, message):
            errors.append(message)
            issues.append(_issue(SHEET_PROPERTIES, row_index, column, message))

        name = _cell_text(raw.get("Property Name"))
        region = _cell_text(raw.get("Region"))
        city = _cell_text(raw.get("City"))
        if action == "create" or not _is_blank(raw.get("Property Name")):
            if not name:
                add_error("Property Name", "Property Name is required.")
        if action == "create" or not _is_blank(raw.get("Region")):
            if not region:
                add_error("Region", "Region is required.")
        if action == "create" or not _is_blank(raw.get("City")):
            if not city:
                add_error("City", "City is required.")

        try:
            if action == "create" or not _is_blank(raw.get("Base Rate")):
                from accounts.property_services import parse_rate

                values["rate"] = str(parse_rate(raw.get("Base Rate")))
        except ValueError as exc:
            add_error("Base Rate", str(exc))

        try:
            if not _is_blank(raw.get("Property Type")) or action == "create":
                values["property_type"] = _parse_choice(
                    raw.get("Property Type"),
                    PropertyType,
                    field="Property Type",
                    default=PropertyType.HOTEL,
                )
        except ValueError as exc:
            add_error("Property Type", str(exc))

        try:
            if not _is_blank(raw.get("Status")) or action == "create":
                values["status"] = _parse_choice(
                    raw.get("Status"),
                    PropertyStatus,
                    field="Status",
                    default=PropertyStatus.ACTIVE,
                )
        except ValueError as exc:
            add_error("Status", str(exc))

        try:
            if not _is_blank(raw.get("Email")):
                from accounts.property_services import parse_optional_email

                values["email"] = parse_optional_email(raw.get("Email"))
        except ValueError as exc:
            add_error("Email", str(exc))

        try:
            if not _is_blank(raw.get("Currency")) or action == "create":
                from accounts.property_services import parse_currency

                values["currency"] = parse_currency(raw.get("Currency") if not _is_blank(raw.get("Currency")) else "MYR")
        except ValueError as exc:
            add_error("Currency", str(exc))

        try:
            if not _is_blank(raw.get("Timezone")) or action == "create":
                from accounts.property_services import parse_timezone

                values["timezone"] = parse_timezone(raw.get("Timezone") if not _is_blank(raw.get("Timezone")) else None)
        except ValueError as exc:
            add_error("Timezone", str(exc))

        try:
            check_in = _parse_time(raw.get("Check-in Time"), field="Check-in Time")
            if check_in:
                values["check_in_time"] = check_in
        except ValueError as exc:
            add_error("Check-in Time", str(exc))

        try:
            check_out = _parse_time(raw.get("Check-out Time"), field="Check-out Time")
            if check_out:
                values["check_out_time"] = check_out
        except ValueError as exc:
            add_error("Check-out Time", str(exc))

        if name:
            values["name"] = name
        if code or action == "create":
            values["code"] = code
        if region:
            values["region"] = region
        if city:
            values["city"] = city
        if not _is_blank(raw.get("Country")) or action == "create":
            values["country"] = _cell_text(raw.get("Country")) or "Malaysia"
        if not _is_blank(raw.get("Address")):
            values["address_line1"] = _cell_text(raw.get("Address"))
        if not _is_blank(raw.get("Address Line 2")):
            values["address_line2"] = _cell_text(raw.get("Address Line 2"))
        if not _is_blank(raw.get("Postal Code")):
            values["postal_code"] = _cell_text(raw.get("Postal Code"))
        if not _is_blank(raw.get("Phone")):
            values["phone"] = _cell_text(raw.get("Phone"))
        if not _is_blank(raw.get("Website")):
            values["website"] = _cell_text(raw.get("Website"))
        if not _is_blank(raw.get("Description")):
            values["description"] = _cell_text(raw.get("Description"))

        results.append(
            {
                "row": row_index,
                "action": action,
                "id": property_obj.pk if property_obj else property_id,
                "name": name or (property_obj.name if property_obj else ""),
                "errors": errors,
                "warnings": warnings,
                "values": values,
            }
        )
    return results


def _validate_location_rows(user, rows, issues, *, property_rows, pending_codes, locked_property_id=None):
    results = []
    seen_ids = {}
    seen_codes = {}
    can_change = user_has_permission(user, "properties.change_property")
    created_codes = {
        (row["values"].get("code") or "").upper(): row
        for row in property_rows
        if row["action"] == "create" and row["values"].get("code")
    }

    for row_index, raw in rows:
        errors = []
        warnings = []
        values = {}
        action = "create"
        location = None
        property_obj = None

        def add_error(column, message):
            errors.append(message)
            issues.append(_issue(SHEET_LOCATIONS, row_index, column, message))

        try:
            location_id = _parse_int(raw.get("Location ID"), field="Location ID")
        except ValueError as exc:
            location_id = None
            add_error("Location ID", str(exc))

        try:
            property_id = _parse_int(raw.get("Property ID"), field="Property ID")
        except ValueError as exc:
            property_id = None
            add_error("Property ID", str(exc))

        property_code = _cell_text(raw.get("Property Code")).upper()
        if locked_property_id:
            if property_id not in (None, locked_property_id):
                add_error("Property ID", "Locations can only be imported into the selected property.")
            property_id = locked_property_id

        if property_id:
            property_obj = Property.objects.filter(pk=property_id).first()
            if property_obj is None:
                add_error("Property ID", "Property ID does not exist.")
            else:
                try:
                    if not can_access_property(user, property_obj):
                        raise PermissionError("You do not have access to this property.")
                except PermissionError as exc:
                    add_error("Property ID", str(exc))
        elif property_code:
            property_obj = Property.objects.filter(code=property_code).first()
            if property_obj is None and property_code in created_codes:
                values["pending_property_code"] = property_code
            elif property_obj is None:
                add_error("Property Code", "Property Code does not match an existing or imported property.")
            else:
                try:
                    if not can_access_property(user, property_obj):
                        raise PermissionError("You do not have access to this property.")
                except PermissionError as exc:
                    add_error("Property Code", str(exc))
                property_id = property_obj.pk
        else:
            add_error("Property ID", "Property ID or Property Code is required.")

        if location_id:
            action = "update"
            if location_id in seen_ids:
                add_error("Location ID", "Duplicate Location ID in this workbook.")
            seen_ids[location_id] = row_index
            location = Location.objects.select_related("property").filter(pk=location_id).first()
            if location is None:
                add_error("Location ID", "Location ID does not exist.")
            else:
                try:
                    if not can_access_property(user, location.property):
                        raise PermissionError("You do not have access to this property.")
                except PermissionError as exc:
                    add_error("Location ID", str(exc))
                if property_obj and location.property_id != property_obj.pk:
                    add_error("Property ID", "A location cannot be moved to another property.")
                elif not property_obj:
                    property_obj = location.property
                    property_id = location.property_id

        code = _cell_text(raw.get("Location Code"))
        if not code:
            add_error("Location Code", "Location Code is required.")
        else:
            key = (property_id or property_code, code.upper())
            if key in seen_codes:
                add_error("Location Code", "Duplicate Location Code for this property in this workbook.")
            seen_codes[key] = row_index
            if property_obj and not location:
                existing = Location.objects.filter(property=property_obj, code=code).first()
                if existing:
                    action = "update"
                    location = existing
                    location_id = existing.pk
                    warnings.append("Matched existing location by Property + Location Code.")
                    issues.append(_issue(SHEET_LOCATIONS, row_index, "Location Code", warnings[-1], "warning"))

        if not can_change:
            add_error("Location Name", "You do not have permission to create or update locations.")

        name = _cell_text(raw.get("Location Name"))
        if not name:
            add_error("Location Name", "Location Name is required.")

        try:
            values["type"] = _parse_choice(raw.get("Location Type"), LocationType, field="Location Type", required=True)
        except ValueError as exc:
            add_error("Location Type", str(exc))

        try:
            values["status"] = _parse_choice(raw.get("Status"), LocationStatus, field="Status", required=True)
        except ValueError as exc:
            add_error("Status", str(exc))

        try:
            values["capacity"] = _parse_int(raw.get("Capacity"), field="Capacity", required=True, minimum=1)
        except ValueError as exc:
            add_error("Capacity", str(exc))

        try:
            if not _is_blank(raw.get("Max Adults")):
                values["max_adults"] = _parse_int(raw.get("Max Adults"), field="Max Adults", minimum=1)
        except ValueError as exc:
            add_error("Max Adults", str(exc))

        try:
            if not _is_blank(raw.get("Max Children")):
                values["max_children"] = _parse_int(raw.get("Max Children"), field="Max Children", minimum=0, allow_zero=True)
        except ValueError as exc:
            add_error("Max Children", str(exc))

        try:
            if not _is_blank(raw.get("Bed Type")):
                values["bed_type"] = _parse_choice(raw.get("Bed Type"), BedType, field="Bed Type")
        except ValueError as exc:
            add_error("Bed Type", str(exc))

        try:
            if not _is_blank(raw.get("Bed Count")):
                values["bed_count"] = _parse_int(raw.get("Bed Count"), field="Bed Count", minimum=1)
        except ValueError as exc:
            add_error("Bed Count", str(exc))

        try:
            if not _is_blank(raw.get("Base Rate")):
                from locations.services import parse_optional_rate

                rate = parse_optional_rate(raw.get("Base Rate"))
                values["base_rate"] = None if rate is None else f"{rate:.2f}"
        except Exception as exc:
            add_error("Base Rate", str(exc))

        try:
            if not _is_blank(raw.get("Currency")):
                from accounts.property_services import parse_currency

                values["currency"] = parse_currency(raw.get("Currency"))
        except ValueError as exc:
            add_error("Currency", str(exc))

        try:
            if not _is_blank(raw.get("Is Bookable")):
                values["is_bookable"] = _parse_bool(raw.get("Is Bookable"), field="Is Bookable", default=True)
        except ValueError as exc:
            add_error("Is Bookable", str(exc))

        if name:
            values["name"] = name
        values["code"] = code
        if not _is_blank(raw.get("Floor")):
            values["floor"] = _cell_text(raw.get("Floor"))
        if not _is_blank(raw.get("Description")):
            values["description"] = _cell_text(raw.get("Description"))
        if property_id:
            values["property_id"] = property_id

        results.append(
            {
                "row": row_index,
                "action": action,
                "id": location.pk if location else location_id,
                "property_id": property_id,
                "name": name,
                "errors": errors,
                "warnings": warnings,
                "values": values,
            }
        )
    return results


@transaction.atomic
def commit_import(user, uploaded, *, mode: str = "portfolio", locked_property_id: int | None = None) -> dict:
    preview = preview_import(user, uploaded, mode=mode, locked_property_id=locked_property_id)
    if not preview["valid"]:
        raise ExcelImportError("The workbook failed validation. Nothing was imported.", code="import_invalid")

    created_properties = []
    updated_properties = []
    created_locations = []
    updated_locations = []
    code_to_property = {}

    for row in preview["properties"]:
        values = dict(row["values"])
        if row["action"] == "create":
            property_obj = create_property(
                actor=user,
                name=values.pop("name"),
                region=values.pop("region"),
                city=values.pop("city"),
                rate=values.pop("rate"),
                **values,
            )
            created_properties.append(property_obj.pk)
            if property_obj.code:
                code_to_property[property_obj.code.upper()] = property_obj
        else:
            property_obj = Property.objects.get(pk=row["id"])
            property_obj = update_property(property_obj, values)
            updated_properties.append(property_obj.pk)
            if property_obj.code:
                code_to_property[property_obj.code.upper()] = property_obj

    for row in preview["locations"]:
        values = dict(row["values"])
        pending_code = values.pop("pending_property_code", None)
        property_id = values.pop("property_id", None) or row.get("property_id")
        if property_id:
            property_obj = Property.objects.get(pk=property_id)
        else:
            property_obj = code_to_property.get((pending_code or "").upper())
            if property_obj is None:
                raise ExcelImportError("A location referenced a property that was not imported.", code="invalid_property")
        payload = values
        if row["action"] == "create":
            location = create_location(property_obj, payload)
            created_locations.append(location.pk)
        else:
            location = Location.objects.get(pk=row["id"])
            if location.property_id != property_obj.pk:
                raise ExcelImportError("A location cannot be moved to another property.", code="location_move_forbidden")
            location = update_location(location, payload)
            updated_locations.append(location.pk)

    return {
        "ok": True,
        "resource": "property_import",
        "mode": mode,
        "created_properties": created_properties,
        "updated_properties": updated_properties,
        "created_locations": created_locations,
        "updated_locations": updated_locations,
    }
